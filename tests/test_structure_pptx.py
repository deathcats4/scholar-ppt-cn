from __future__ import annotations

import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches

from scripts.structure_pptx import NS, StructurePlanError, structure_pptx


P = f"{{{NS['p']}}}"
A = f"{{{NS['a']}}}"


def make_fixture(path: Path, *, duplicate_name: bool = False, rotated: bool = False) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])

    node_a = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1), Inches(1), Inches(0.5), Inches(0.5))
    node_a.name = "NODE_A"
    if rotated:
        node_a.rotation = 15
    label_a = slide.shapes.add_textbox(Inches(1.55), Inches(1), Inches(1), Inches(0.5))
    label_a.name = "NODE_A" if duplicate_name else "LABEL_A"
    label_a.text = "A"

    node_b = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4), Inches(1), Inches(0.5), Inches(0.5))
    node_b.name = "NODE_B"
    label_b = slide.shapes.add_textbox(Inches(4.55), Inches(1), Inches(1), Inches(0.5))
    label_b.name = "LABEL_B"
    label_b.text = "B"

    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(2.7),
        Inches(1.25),
        Inches(3.8),
        Inches(1.25),
    )
    connector.name = "LINE_A_B"

    untouched = deck.slides.add_slide(deck.slide_layouts[6])
    marker = untouched.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    marker.name = "UNTOUCHED"
    marker.text = "unchanged package part"
    deck.save(path)


def package_bytes(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as package:
        return {name: package.read(name) for name in package.namelist()}


def replace_package_part(path: Path, part_name: str, data: bytes) -> None:
    temporary = path.with_suffix(".rewrite.tmp")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w") as destination:
        destination.comment = source.comment
        for info in source.infolist():
            destination.writestr(info, data if info.filename == part_name else source.read(info.filename))
    os.replace(temporary, path)


def cnvpr(shape: ET.Element) -> ET.Element:
    properties = shape.find(".//p:cNvPr", NS)
    if properties is None:
        raise AssertionError("shape lacks cNvPr")
    return properties


def shape_by_name(root: ET.Element, name: str) -> ET.Element:
    for shape in root.iter():
        if shape.tag in {P + item for item in ("sp", "pic", "graphicFrame", "cxnSp", "grpSp")}:
            properties = shape.find(".//p:cNvPr", NS)
            if properties is not None and properties.attrib.get("name") == name:
                return shape
    raise AssertionError(f"shape {name!r} was not found")


def xfrm_bytes(shape: ET.Element) -> bytes:
    transform = shape.find("p:spPr/a:xfrm", NS)
    if transform is None:
        transform = shape.find("p:xfrm", NS)
    if transform is None:
        raise AssertionError("shape lacks xfrm")
    return ET.tostring(transform)


def convert_line_to_plain_shape(path: Path) -> None:
    parts = package_bytes(path)
    root = ET.fromstring(parts["ppt/slides/slide1.xml"])
    connector = shape_by_name(root, "LINE_A_B")
    if connector.tag != P + "cxnSp":
        raise AssertionError("python-pptx did not create a connector shape")
    connector.tag = P + "sp"
    non_visual = connector.find("p:nvCxnSpPr", NS)
    if non_visual is None:
        raise AssertionError("connector lacks nvCxnSpPr")
    non_visual.tag = P + "nvSpPr"
    connection_properties = non_visual.find("p:cNvCxnSpPr", NS)
    if connection_properties is None:
        raise AssertionError("connector lacks cNvCxnSpPr")
    connection_properties.tag = P + "cNvSpPr"
    replace_package_part(path, "ppt/slides/slide1.xml", ET.tostring(root, xml_declaration=True))


def add_ignorable_namespace_declarations(path: Path) -> None:
    parts = package_bytes(path)
    data = parts["ppt/slides/slide1.xml"]
    marker = b"<p:sld "
    replacement = (
        b'<p:sld xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        b'xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" '
        b'mc:Ignorable="a14" '
    )
    if marker not in data:
        raise AssertionError("unexpected slide root prefix")
    replace_package_part(path, "ppt/slides/slide1.xml", data.replace(marker, replacement, 1))


def valid_plan() -> dict[str, object]:
    return {
        "groups": [
            {"slide": 1, "name": "GROUP_A", "members": ["NODE_A", "LABEL_A"]},
            {"slide": 1, "name": "GROUP_B", "members": ["NODE_B", "LABEL_B"]},
        ],
        "connectors": [
            {
                "slide": 1,
                "name": "LINE_A_B",
                "from": {"shape": "NODE_A", "site": 2},
                "to": {"shape": "NODE_B", "site": 0},
            }
        ],
    }


class StructurePptxTests(unittest.TestCase):
    def test_groups_and_connector_preserve_geometry_ids_and_other_parts(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root_dir = Path(tmp)
            source = root_dir / "source.pptx"
            output = root_dir / "structured.pptx"
            make_fixture(source)
            convert_line_to_plain_shape(source)
            add_ignorable_namespace_declarations(source)
            before = package_bytes(source)
            before_root = ET.fromstring(before["ppt/slides/slide1.xml"])
            original_geometry = {
                name: xfrm_bytes(shape_by_name(before_root, name))
                for name in ("NODE_A", "LABEL_A", "NODE_B", "LABEL_B", "LINE_A_B")
            }
            original_ids = {
                name: cnvpr(shape_by_name(before_root, name)).attrib["id"]
                for name in ("NODE_A", "LABEL_A", "NODE_B", "LABEL_B", "LINE_A_B")
            }

            report = structure_pptx(source, valid_plan(), output)

            self.assertTrue(output.is_file())
            reopened = Presentation(output)
            self.assertEqual(len(reopened.slides), 2)
            self.assertEqual(report["status"], "passed")
            self.assertEqual(report["details"]["groups_created"], 2)
            self.assertEqual(report["details"]["connectors_bound"], 1)
            self.assertEqual(report["details"]["plain_lines_converted"], 1)

            after = package_bytes(output)
            self.assertEqual(set(after), set(before))
            for part_name, data in before.items():
                if part_name != "ppt/slides/slide1.xml":
                    self.assertEqual(after[part_name], data, part_name)

            after_root = ET.fromstring(after["ppt/slides/slide1.xml"])
            ids = [int(node.attrib["id"]) for node in after_root.findall(".//p:cNvPr", NS)]
            self.assertEqual(len(ids), len(set(ids)))
            for name, expected_id in original_ids.items():
                self.assertEqual(cnvpr(shape_by_name(after_root, name)).attrib["id"], expected_id)
            for name, expected_xfrm in original_geometry.items():
                self.assertEqual(xfrm_bytes(shape_by_name(after_root, name)), expected_xfrm)

            group_a = shape_by_name(after_root, "GROUP_A")
            group_member_names = [
                cnvpr(child).attrib["name"]
                for child in list(group_a)
                if child.tag in {P + "sp", P + "pic", P + "graphicFrame", P + "cxnSp"}
            ]
            self.assertEqual(group_member_names, ["NODE_A", "LABEL_A"])
            group_xfrm = group_a.find("p:grpSpPr/a:xfrm", NS)
            self.assertIsNotNone(group_xfrm)
            self.assertEqual(group_xfrm.find("a:off", NS).attrib, group_xfrm.find("a:chOff", NS).attrib)
            self.assertEqual(group_xfrm.find("a:ext", NS).attrib, group_xfrm.find("a:chExt", NS).attrib)

            connector = shape_by_name(after_root, "LINE_A_B")
            self.assertEqual(connector.tag, P + "cxnSp")
            start = connector.find("p:nvCxnSpPr/p:cNvCxnSpPr/a:stCxn", NS)
            end = connector.find("p:nvCxnSpPr/p:cNvCxnSpPr/a:endCxn", NS)
            self.assertEqual(start.attrib, {"id": original_ids["NODE_A"], "idx": "2"})
            self.assertEqual(end.attrib, {"id": original_ids["NODE_B"], "idx": "0"})
            slide_xml = after["ppt/slides/slide1.xml"]
            self.assertIn(b'mc:Ignorable="a14"', slide_xml)
            self.assertIn(b'xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main"', slide_xml)

    def test_invalid_plans_do_not_create_output(self) -> None:
        cases = {
            "illegal-slide": {
                "groups": [{"slide": 3, "name": "G", "members": ["NODE_A", "LABEL_A"]}]
            },
            "missing-name": {
                "groups": [{"slide": 1, "name": "G", "members": ["NODE_A", "MISSING"]}]
            },
            "duplicate-member": {
                "groups": [{"slide": 1, "name": "G", "members": ["NODE_A", "NODE_A"]}]
            },
            "non-contiguous": {
                "groups": [{"slide": 1, "name": "G", "members": ["NODE_A", "NODE_B"]}]
            },
            "bad-site": {
                "connectors": [
                    {
                        "slide": 1,
                        "name": "LINE_A_B",
                        "from": {"shape": "NODE_A", "site": "right"},
                        "to": {"shape": "NODE_B", "site": 0},
                    }
                ]
            },
        }
        with tempfile.TemporaryDirectory() as tmp:
            root_dir = Path(tmp)
            source = root_dir / "source.pptx"
            make_fixture(source)
            source_bytes = source.read_bytes()
            for case_name, plan in cases.items():
                output = root_dir / f"{case_name}.pptx"
                with self.subTest(case=case_name):
                    with self.assertRaises(StructurePlanError):
                        structure_pptx(source, plan, output)
                    self.assertFalse(output.exists())
                    self.assertEqual(source.read_bytes(), source_bytes)

    def test_ambiguous_name_and_rotated_group_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root_dir = Path(tmp)
            ambiguous = root_dir / "ambiguous.pptx"
            make_fixture(ambiguous, duplicate_name=True)
            with self.assertRaisesRegex(StructurePlanError, "not unique"):
                structure_pptx(
                    ambiguous,
                    {"groups": [{"slide": 1, "name": "G", "members": ["NODE_A", "NODE_B"]}]},
                    root_dir / "ambiguous-output.pptx",
                )

            rotated = root_dir / "rotated.pptx"
            make_fixture(rotated, rotated=True)
            with self.assertRaisesRegex(StructurePlanError, "Rotated or flipped"):
                structure_pptx(
                    rotated,
                    {"groups": [{"slide": 1, "name": "G", "members": ["NODE_A", "LABEL_A"]}]},
                    root_dir / "rotated-output.pptx",
                )
            self.assertFalse((root_dir / "ambiguous-output.pptx").exists())
            self.assertFalse((root_dir / "rotated-output.pptx").exists())

    def test_namespace_prefix_rebinding_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.pptx"
            output = Path(tmp) / "output.pptx"
            make_fixture(source)
            parts = package_bytes(source)
            xml = parts["ppt/slides/slide1.xml"]
            xml = xml.replace(b"<p:sld ", b'<p:sld xmlns:foo="urn:outer" ', 1)
            xml = xml.replace(b"<p:cSld>", b'<p:cSld xmlns:foo="urn:inner">', 1)
            self.assertIn(b'xmlns:foo="urn:outer"', xml)
            self.assertIn(b'xmlns:foo="urn:inner"', xml)
            replace_package_part(source, "ppt/slides/slide1.xml", xml)
            before = source.read_bytes()
            with self.assertRaisesRegex(StructurePlanError, "rebound"):
                structure_pptx(source, {}, output)
            self.assertFalse(output.exists())
            self.assertEqual(source.read_bytes(), before)

    def test_plain_line_text_is_rejected_without_losing_content(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.pptx"
            output = Path(tmp) / "output.pptx"
            make_fixture(source)
            convert_line_to_plain_shape(source)
            parts = package_bytes(source)
            root = ET.fromstring(parts["ppt/slides/slide1.xml"])
            line = shape_by_name(root, "LINE_A_B")
            text_body = ET.SubElement(line, P + "txBody")
            ET.SubElement(text_body, A + "bodyPr")
            ET.SubElement(text_body, A + "lstStyle")
            paragraph = ET.SubElement(text_body, A + "p")
            run = ET.SubElement(paragraph, A + "r")
            ET.SubElement(run, A + "t").text = "Keep this label"
            replace_package_part(source, "ppt/slides/slide1.xml", ET.tostring(root))
            before = source.read_bytes()
            plan = {"connectors": [{
                "slide": 1, "name": "LINE_A_B",
                "from": {"shape": "NODE_A", "site": 6},
                "to": {"shape": "NODE_B", "site": 2},
            }]}
            with self.assertRaisesRegex(StructurePlanError, "cannot preserve"):
                structure_pptx(source, plan, output)
            self.assertFalse(output.exists())
            self.assertEqual(source.read_bytes(), before)

    def test_plain_line_cannot_connect_to_itself(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.pptx"
            output = Path(tmp) / "output.pptx"
            make_fixture(source)
            convert_line_to_plain_shape(source)
            before = source.read_bytes()
            for self_endpoint in ("from", "to"):
                connector = {
                    "slide": 1, "name": "LINE_A_B",
                    "from": {"shape": "NODE_A", "site": 0},
                    "to": {"shape": "NODE_B", "site": 0},
                }
                connector[self_endpoint]["shape"] = "LINE_A_B"
                with self.subTest(endpoint=self_endpoint):
                    with self.assertRaisesRegex(StructurePlanError, "connect to itself"):
                        structure_pptx(source, {"connectors": [connector]}, output)
                    self.assertFalse(output.exists())
                    self.assertEqual(source.read_bytes(), before)

    def test_in_place_output_is_refused_without_changing_input(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.pptx"
            make_fixture(source)
            before = source.read_bytes()

            with self.assertRaisesRegex(StructurePlanError, "in-place overwrite"):
                structure_pptx(source, {}, source)

            self.assertEqual(source.read_bytes(), before)
            self.assertGreater(len(Presentation(source).slides), 0)


if __name__ == "__main__":
    unittest.main()
